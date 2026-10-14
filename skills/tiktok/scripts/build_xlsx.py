#!/usr/bin/env python3
"""把调研结果生成带图片的 xlsx。
输入 CSV 或 JSON(list)，字段: sku,name,cost,case,image_path,comp_price,sold,result,link
用法: build_xlsx.py rows.csv out.xlsx
图片用 TwoCellAnchor(editAs='oneCell') 嵌入 A 列（WPS/Excel 都能显示；OneCellAnchor 在 WPS 不显示）。"""
import csv, json, sys, io
from pathlib import Path
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.styles import Font, Alignment, PatternFill
from PIL import Image as PILImage

if len(sys.argv)<3 or sys.argv[1] in ('-h','--help'):
    print(__doc__); sys.exit(0)
src, out = sys.argv[1], sys.argv[2]
rows = json.load(open(src)) if src.endswith('.json') else list(csv.DictReader(open(src, encoding='utf-8-sig')))
HEAD = ['图片','SKU','商品','成本价(R$)','整箱','售价','保本价','利润','竞品(R$)','销量','结果','链接']
KEYS = [None,'sku','name','cost','case',None,None,None,'comp_price','sold','result','link']
WID = [22,14,30,11,8,9,9,9,11,10,40,45]
wb = Workbook(); ws = wb.active; ws.title = '调研'
for i,(h,w) in enumerate(zip(HEAD,WID), 1):
    c = ws.cell(1, i, h); c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBF7')
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[chr(64+i)].width = w
def num(v):
    try: return float(v) if '.' in str(v) else int(v)
    except Exception: return v if v not in (None,) else ''
for r, row in enumerate(rows, 2):
    ws.row_dimensions[r].height = 110
    for i,k in enumerate(KEYS, 1):
        if not k: continue
        v = row.get(k, '')
        v = num(v) if k in ('cost','case','comp_price','sold') and v not in ('', None) else v
        c = ws.cell(r, i, v); c.alignment = Alignment(vertical='center', wrap_text=True)
        if k == 'link' and v: c.hyperlink = v; c.font = Font(color='0563C1', underline='single')
        if k == 'result' and v and '【' in str(v): c.font = Font(color='C00000', bold=True)
    p = row.get('image_path')
    if p and Path(p).exists():
        im = PILImage.open(p).convert('RGB'); im.thumbnail((300, 300))
        buf = io.BytesIO(); im.save(buf, 'JPEG', quality=85); buf.seek(0)
        xi = XLImage(buf)
        a = TwoCellAnchor(editAs='oneCell')
        a._from = AnchorMarker(col=0, row=r-1, colOff=0, rowOff=0)
        a.to = AnchorMarker(col=1, row=r, colOff=0, rowOff=0)
        xi.anchor = a; ws.add_image(xi)
ws.freeze_panes = 'B2'
wb.save(out); print(f'已生成 {out}，{len(rows)} 行')
